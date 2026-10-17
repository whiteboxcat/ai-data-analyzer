"""Workbook-level cleaning: lookup sheets, reruns, side pivots, notes sheets.

Builds a small workbook with the same patterns as a typical BI data bank:
    python tests/test_workbook.py
"""
import datetime as dt
import sys
import tempfile
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from analyzer import industry  # noqa: E402
from analyzer.cleaner import detect_issues, loose_key  # noqa: E402
from analyzer.loader import load_file  # noqa: E402
from analyzer.pipeline import run  # noqa: E402
from analyzer.profiler import profile_sheet  # noqa: E402


def build(path: Path, with_masters: bool = True) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    ws.append(["Data bank (synthetic)"])
    ws.append(["Not every mismatch is an error: some are caused by timing or business definition."])

    if with_masters:
        pm = wb.create_sheet("product_master")
        pm.append(["sku_raw", "sku_canonical", "product_name", "unit_cogs"])
        for raw, canon, name, cogs in [("FBL 200ml", "FBL200", "Lotion 200ml", 47500), ("FBL-200", "FBL200", "Lotion 200ml", 51000),
                                       ("HTTW450", "HTTW450", "Wash 450ml", 43000), ("TELON 110", "TELON110", "Telon oil", 34000)]:
            pm.append([raw, canon, name, cogs])
        dm = wb.create_sheet("distributor_master")
        dm.append(["distributor_raw", "distributor_canonical", "region"])
        for raw, canon, reg in [("PT Sentral Mapan", "SENTRAL_MAPAN", "Medan"), ("Sentral Mapan", "SENTRAL_MAPAN", "Medan"),
                                ("Maju Bersama", "MAJU_BERSAMA", "Jakarta"), ("Bali Jaya", "BALI_JAYA", "Denpasar")]:
            dm.append([raw, canon, reg])

    si = wb.create_sheet("offline_sell_in")
    si.append(["month", "invoice_id", "distributor_raw", "sku_raw", "units", "gross_sales", "net_sales",
               "extract_version", None, "extract_version", "BASE"])               # pivot filter beside the data
    rows = []
    for i in range(12):
        dist = ["PT Sentral Mapan", "Maju Bersama", "Bali Jaya"][i % 3] if i < 8 else ["Sentral Mapan", "Maju Bersama", "Bali Jaya"][i % 3]
        sku = ["FBL 200ml", "HTTW450", "TELON 110"][i % 3] if i < 8 else ["FBL-200", "HTTW450", "TELON 110"][i % 3]
        rows.append([dt.datetime(2026, 6 + i // 4, 1), f"INV{1001 + i}", dist, sku, 100 + i, 1_000_000 + i * 1000,
                     900_000 + i * 900, "BASE"])
    rows.append(rows[9][:7] + ["RERUN"])                                            # same invoice loaded again
    pivot = [[None, None], ["Row Labels", "Sum of gross_sales"], ["Jun", 4006000], ["Jul", 4022000]]
    for k, r in enumerate(rows):
        si.append(r + [None] + (pivot[k] if k < len(pivot) else [None, None]))

    dd = wb.create_sheet("data_dictionary")
    dd.append(["table", "field", "provided_description"])
    dd.append(["offline_sell_in", "grain", "invoice_id + extract_version; duplicate reruns may exist"])
    wb.save(path)


def analyse(path):
    with mock.patch.object(industry.requests, "get", side_effect=OSError("offline")):
        return run(path, with_trends=False)


def test_with_lookup_sheets(tmp: Path):
    p = tmp / "bank.xlsx"
    build(p)
    sheets = load_file(p)
    assert sheets["README"].attrs.get("doc_text") and sheets["data_dictionary"].attrs.get("doc_text")
    sell = sheets["offline_sell_in"]
    assert list(sell.columns)[-1] == "extract_version" and sell.attrs["side_tables"][0]["columns"] == "J–K"

    r = analyse(p)
    assert r["notes_sheets"] == ["README", "data_dictionary"] and "duplicate reruns" in r["workbook_notes"]
    assert set(r["reference_sheets"]) == {"product_master", "distributor_master"}
    s = next(x for x in r["sheets"] if x["sheet"] == "offline_sell_in")
    kinds = {i["kind"]: i for i in s["issues"]}
    sku = next(i for i in s["issues"] if i["kind"] == "mapped_names" and i["column"] == "sku_raw")
    assert "'FBL 200ml', 'FBL-200' → FBL200" in sku["detail"], sku["detail"]
    dist = next(i for i in s["issues"] if i["kind"] == "mapped_names" and i["column"] == "distributor_raw")
    assert "SENTRAL_MAPAN" in dist["detail"]
    assert "INV1010" in kinds["reloaded_rows"]["detail"] and "RERUN" in kinds["reloaded_rows"]["detail"]
    assert kinds["side_table"]["info"] is True
    assert "duplicate_ids" not in kinds and "similar_labels" not in kinds
    assert s["rows"] == 12                                                          # rerun dropped
    cols = [c["name"] for c in s["profile"]["columns"]]
    assert cols.index("sku_canonical") == cols.index("sku_raw") + 1
    assert not any("_raw" in (c["x"] or "") for c in s["charts"]), [c["x"] for c in s["charts"]]
    pm = next(x for x in r["sheets"] if x["sheet"] == "product_master")
    assert pm["charts"] == [] and "Lookup table" in pm["description"]
    print("OK lookup sheets: FBL 200ml/FBL-200 -> FBL200, Sentral Mapan merged, rerun dropped, pivot ignored, notes read")


def test_without_lookup_sheets(tmp: Path):
    p = tmp / "plain.xlsx"
    build(p, with_masters=False)
    df = load_file(p)["offline_sell_in"]
    issues = detect_issues(df, profile_sheet(df, "offline_sell_in"))
    similar = [i for i in issues if i["kind"] == "similar_labels"]
    assert {i["column"] for i in similar} == {"sku_raw", "distributor_raw"}, [i["title"] for i in similar]
    sku = next(i for i in similar if i["column"] == "sku_raw")
    assert "'FBL 200ml' / 'FBL-200'" in sku["detail"] and not sku["auto_fix"]
    assert loose_key("PT. Sentral Mapan") == loose_key("SentralMapan") == "sentralmapan"
    assert loose_key("FBL 200ml") == loose_key("FBL-200") == "fbl200" != loose_key("FBL400")
    print("OK without a master: similar spellings flagged for review (not merged automatically)")


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as d:
        test_with_lookup_sheets(Path(d))
        test_without_lookup_sheets(Path(d))
